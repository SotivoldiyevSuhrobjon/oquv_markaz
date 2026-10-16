import datetime as dt

import openpyxl
import pytest
from django.core.management import call_command

from apps.academics.models import Group, Student
from apps.accounts.models import User


def pad(row, n=14):
    return row + [None] * (n - len(row))


def build_workbook(path):
    wb = openpyxl.Workbook()
    wb.active.title = "Kids"
    for name in ("Matematika", "English", "Xonalar"):
        wb.create_sheet(name)
    ws = wb["Matematika"]
    ws.append(pad(["Mat_1", "Rahimjon", None, None, None, None, None, dt.time(8, 0), None, "101"]))
    ws.append(pad([None, None, dt.datetime(2026, 9, 3), "Y", "Ali Valiyev", "+998 90 123 45 67", "D", None, "7"]))
    ws.append(pad([None, None, None, "N", "Yo'q Bola", None, "D"]))
    u = wb.create_sheet("Ustozlar")
    u.append(pad(["Matematika fani"]))
    u.append(pad([None, "F.I.Sh", "Tel"]))
    u.append(pad([None, "Rahimjon", "998901112233"]))
    wb.save(path)


@pytest.mark.django_db
def test_import_creates_groups_students_staff(tmp_path):
    f = tmp_path / "g.xlsx"
    build_workbook(f)
    call_command("import_excel", str(f), "--password-for-staff", "Zor-Parol-9876")
    assert Group.objects.filter(name="Mat_1", room="101", days="DCJ").exists()
    assert Student.objects.get(full_name="Ali Valiyev").enrollments.count() == 1
    assert not Student.objects.filter(full_name="Yo'q Bola").exists()
    assert User.objects.get(username="rahimjon").must_change_password


@pytest.mark.django_db
def test_dry_run_saves_nothing(tmp_path):
    f = tmp_path / "g.xlsx"
    build_workbook(f)
    call_command("import_excel", str(f), "--dry-run")
    assert Student.objects.count() == 0


def test_qabul_dates():
    from apps.academics.management.commands.import_excel import phone_of, qabul_of

    assert qabul_of("21.sent") == dt.date(2026, 9, 21)
    assert qabul_of("15,10") == dt.date(2026, 10, 15)
    assert qabul_of("okt") == dt.date(2026, 10, 1)
    assert qabul_of(None) == dt.date(2026, 9, 1)
    assert phone_of(883723377) == "+998883723377"


@pytest.mark.django_db
def test_init_center_sets_admin_for_site_and_django_admin(client):
    from django.core.management import call_command

    call_command("init_center", "--password", "Admin123.")
    assert client.login(username="admin", password="Admin123.")
    assert client.get("/admin/").status_code == 200  # Django admin ham shu login bilan
    call_command("init_center", "--password", "Yangi-Parol-55")
    assert User.objects.get(username="admin").check_password("Yangi-Parol-55")
